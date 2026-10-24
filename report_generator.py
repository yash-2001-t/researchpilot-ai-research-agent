
from io import BytesIO
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


WORDS_PER_PAGE = 400


def target_word_count(page_count: int) -> int:
    """Estimate a report's target word count from its page setting."""
    if page_count not in (2, 5, 10, 15):
        raise ValueError("Page count must be 2, 5, 10, or 15.")

    return page_count * WORDS_PER_PAGE


def build_report_prompt(
    question: str,
    evidence: str,
    output_mode: str,
    page_count: int,
) -> str:
    """Build the prompt for an abstract or full research report."""
    if output_mode not in ("abstract", "full_report"):
        raise ValueError("Unsupported report output mode.")

    if output_mode == "abstract":
        output_instructions = """
Write a standalone research abstract of approximately 200–300 words.

Include:
- The research topic and scope.
- The most important findings supported by the sources.
- Relevant disagreement, limitations, or uncertainty.
- A concise overall conclusion.

Do not write a full report with multiple long sections.
"""
    else:
        words = target_word_count(page_count)

        output_instructions = f"""
Write a substantial research report targeting approximately {words}
words, equivalent to about {page_count} Word pages under standard
formatting. This is an approximate target, not an exact page guarantee.

Use this structure:
1. Title
2. Executive summary
3. Introduction and scope
4. Thematic analysis and key findings
5. Contradictory findings and alternative explanations
6. Limitations and evidence gaps
7. Implications and conclusion

Adjust the depth of analysis to the selected length. Avoid repetition
and do not pad the report merely to meet a word target.
"""

    return f"""
You are ResearchPilot, an evidence-based research assistant.

Research question or topic:
{question}

Evidence retrieved for this research:
{evidence}

{output_instructions}

Citation and evidence requirements:
- Use Markdown headings and readable paragraphs.
- Cite factual claims with links to the supplied source URLs.
- Use only the supplied source URLs; do not invent citations,
  authors, publication titles, statistics, or quotations.
- Distinguish findings from interpretations.
- Explain important disagreements and limitations.
- Do not claim to have read a complete academic paper when only
  its abstract or metadata was supplied.
- Acknowledge when the available evidence is insufficient.
- End with a concise conclusion appropriate to the output mode.
"""


def build_markdown_export(report: str, sources: list[dict]) -> str:
    """Append a source bibliography to a Markdown report."""
    lines = [
        report.strip(),
        "",
        "---",
        "",
        "## Sources",
        "",
    ]

    for index, source in enumerate(sources, start=1):
        title = source.get("title") or "Untitled source"
        url = source.get("url") or ""
        publication_date = (
            source.get("publication_date") or "Date not verified"
        )
        source_type = source.get("source_type", "web").title()

        if url:
            lines.append(
                f"{index}. [{title}]({url}) — "
                f"{source_type}; publication date: {publication_date}"
            )
        else:
            lines.append(
                f"{index}. {title} — "
                f"{source_type}; publication date: {publication_date}"
            )

    return "\n".join(lines) + "\n"


def _add_inline_markdown(paragraph, text: str):
    """Render basic bold and inline-code formatting in a Word paragraph."""
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    parts = pattern.split(text)

    for part in parts:
        if not part:
            continue

        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def _add_report_body(document, markdown_text: str):
    """Convert common Markdown headings, lists, and paragraphs to Word."""
    for raw_line in markdown_text.splitlines():
        line = raw_line.strip()

        if not line:
            continue

        if line == "---":
            document.add_paragraph()
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.+)$", line)

        if heading_match:
            level = min(len(heading_match.group(1)), 3)
            heading_text = heading_match.group(2).strip()

            heading = document.add_heading(
                heading_text,
                level=level,
            )

            if level == 1:
                heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

            continue

        if line.startswith("- ") or line.startswith("* "):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_markdown(paragraph, line[2:].strip())
            continue

        numbered_match = re.match(r"^\d+[.)]\s+(.+)$", line)

        if numbered_match:
            paragraph = document.add_paragraph(style="List Number")
            _add_inline_markdown(paragraph, numbered_match.group(1))
            continue

        paragraph = document.add_paragraph()
        _add_inline_markdown(paragraph, line)


def build_word_document(
    report: str,
    sources: list[dict],
    title: str = "ResearchPilot Research Report",
) -> bytes:
    """Create a formatted Word document and return its bytes."""
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10.5)
    normal_style.paragraph_format.space_after = Pt(7)

    document.add_heading(title, level=0)

    metadata = document.add_paragraph(
        "Generated by ResearchPilot"
    )
    metadata.alignment = WD_ALIGN_PARAGRAPH.LEFT

    _add_report_body(document, report)

    document.add_page_break()
    document.add_heading("Sources", level=1)

    for index, source in enumerate(sources, start=1):
        source_title = source.get("title") or "Untitled source"
        url = source.get("url") or ""
        source_type = source.get("source_type", "web").title()
        publication_date = (
            source.get("publication_date") or "Date not verified"
        )

        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(source_title).bold = True

        paragraph.add_run(
            f"\nType: {source_type}"
            f"\nPublication date: {publication_date}"
        )

        if url:
            paragraph.add_run(f"\nURL: {url}")

    output = BytesIO()
    document.save(output)

    return output.getvalue()
